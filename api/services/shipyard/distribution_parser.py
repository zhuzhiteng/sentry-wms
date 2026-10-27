"""Parser for the shipyard tray distribution-demand workbook."""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any
from zipfile import BadZipFile
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from services.shipyard.xlsx_security import XlsxSecurityError, validate_xlsx_archive

SOURCE_TIMEZONE = ZoneInfo("Asia/Shanghai")
MAX_SCAN_ROWS = 5000
MAX_TEMPLATE_COLUMNS = 64
HEADER_SEARCH_ROWS = 20
PARSER_VERSION = "shipyard-distribution-v1"

HEADER_ALIASES = {
    "engineering_no": ("工程号", "项目号"),
    "section_no": ("分段号", "分段"),
    "pallet_no": ("托盘号", "托盘编号"),
    "pallet_name": ("托盘名称", "托盘表名称"),
    "work_order_status": ("工单状态",),
    "applicant_raw": ("申请人",),
    "department": ("部门", "申请部门"),
    "delivery_address": ("配送地址", "送货地址", "需求位置"),
    "contact_name": ("联系人",),
    "contact_phone": ("手机号", "联系电话", "电话"),
    "demand_window_raw": ("需求时间", "需求时段"),
    "resource_1": ("配送资源1",),
    "driver_1": ("配送员1",),
    "resource_2": ("配送资源2",),
    "driver_2": ("配送员2",),
    "resource_3": ("配送资源3",),
    "driver_3": ("配送员3",),
    "confirmer": ("确认人",),
    "completed_at_raw": ("完成时间",),
    "confirmation_status": ("确认状态",),
}


@dataclass(frozen=True)
class ParsedDistributionLine:
    source_row_no: str
    source_row_hash: str
    engineering_no: str
    section_no: str
    pallet_no: str
    pallet_name: str
    work_order_status: str | None
    delivery_status: str | None
    applicant_raw: str | None
    department: str | None
    delivery_address: str | None
    contact_name: str | None
    contact_phone: str | None
    demand_window_raw: str | None
    demand_start_at: datetime | None
    demand_end_at: datetime | None
    resource_1: str | None
    driver_1: str | None
    resource_2: str | None
    driver_2: str | None
    resource_3: str | None
    driver_3: str | None
    confirmer: str | None
    completed_at_raw: str | None
    completed_at: datetime | None
    confirmation_status: str | None
    raw_data: dict[str, Any]


@dataclass(frozen=True)
class ParsedDistributionWorkbook:
    sheet_name: str
    header_row: int
    metadata: dict[str, Any]
    warnings: tuple[dict[str, Any], ...]
    lines: tuple[ParsedDistributionLine, ...]


class DistributionParseError(ValueError):
    def __init__(
        self,
        message: str,
        diagnostics: list[dict[str, Any]] | None = None,
    ):
        super().__init__(message)
        self.diagnostics = diagnostics or []


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    return re.sub(r"[\s\u3000:：()（）]+", "", text)


_NORMALIZED_ALIASES = {
    canonical: frozenset(_normalize_header(alias) for alias in aliases)
    for canonical, aliases in HEADER_ALIASES.items()
}


def _to_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        text = value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = text.strip()
    return text or None


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _read_sheet_rows(worksheet) -> list[tuple[int, tuple[Any, ...]]]:
    max_row = min(worksheet.max_row or 0, MAX_SCAN_ROWS)
    if max_row <= 0:
        return []
    max_column = min(worksheet.max_column or 1, MAX_TEMPLATE_COLUMNS)
    return [
        (excel_row, tuple(values))
        for excel_row, values in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=max_row,
                max_col=max_column,
                values_only=True,
            ),
            start=1,
        )
    ]


def _find_header(
    rows: list[tuple[int, tuple[Any, ...]]],
) -> tuple[int, int, dict[str, int], dict[str, str]] | None:
    best = None
    best_score = -1
    for row_offset, (excel_row, values) in enumerate(rows[:HEADER_SEARCH_ROWS]):
        mapping: dict[str, int] = {}
        labels: dict[str, str] = {}
        for column_index, value in enumerate(values[:MAX_TEMPLATE_COLUMNS]):
            normalized = _normalize_header(value)
            if not normalized:
                continue
            for canonical, aliases in _NORMALIZED_ALIASES.items():
                if canonical not in mapping and normalized in aliases:
                    mapping[canonical] = column_index
                    labels[canonical] = str(value).strip()
                    break
        required = {"engineering_no", "section_no", "pallet_no", "pallet_name"}
        if not required.issubset(mapping):
            continue
        score = len(mapping)
        if score > best_score:
            best_score = score
            best = (row_offset, excel_row, mapping, labels)
    return best


def _parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=SOURCE_TIMEZONE) if value.tzinfo is None else value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time(), tzinfo=SOURCE_TIMEZONE)
    text = _to_text(value)
    if not text:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M",
    ):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.replace(tzinfo=SOURCE_TIMEZONE)
        except ValueError:
            continue
    return None


def _parse_window(value: Any) -> tuple[datetime | None, datetime | None]:
    if isinstance(value, datetime):
        return _parse_datetime(value), None
    text = _to_text(value)
    if not text:
        return None, None
    parts = re.split(r"\s*[~～至]\s*", text, maxsplit=1)
    if len(parts) != 2:
        return None, None
    return _parse_datetime(parts[0]), _parse_datetime(parts[1])


def _row_hash(payload: dict[str, Any]) -> str:
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _delivery_status_column(
    rows: list[tuple[int, tuple[Any, ...]]],
    *,
    header_offset: int,
    mapping: dict[str, int],
) -> tuple[int | None, list[dict[str, Any]]]:
    if "delivery_status" in mapping:
        return mapping["delivery_status"], []
    header_values = rows[header_offset][1]
    blank_columns = []
    for column_index, header in enumerate(header_values):
        if _normalize_header(header):
            continue
        has_values = any(
            column_index < len(values) and _to_text(values[column_index])
            for _, values in rows[header_offset + 1 :]
        )
        if has_values:
            blank_columns.append(column_index)
    if len(blank_columns) == 1:
        column_index = blank_columns[0]
        letter = get_column_letter(column_index + 1)
        return column_index, [
            {
                "code": "UNNAMED_HEADER_MAPPED",
                "column": letter,
                "field": "delivery_status",
                "message": f"列 {letter} 没有表头，已按配送状态暂存，需业务确认",
            }
        ]
    if blank_columns:
        return None, [
            {
                "code": "UNNAMED_HEADER_COLUMNS",
                "columns": [get_column_letter(index + 1) for index in blank_columns],
                "message": "存在多个无表头但有数据的列，未自动推断业务含义",
            }
        ]
    return None, []


def _business_rows(
    *,
    rows: list[tuple[int, tuple[Any, ...]]],
    header_offset: int,
    mapping: dict[str, int],
    labels: dict[str, str],
    delivery_status_index: int | None,
    sheet_name: str,
) -> tuple[list[ParsedDistributionLine], list[dict[str, Any]]]:
    lines = []
    diagnostics = []
    seen_pallets = set()
    for excel_row, values in rows[header_offset + 1 :]:
        if not any(_to_text(value) for value in values):
            continue
        fields = {
            field: _to_text(values[index]) if index < len(values) else None
            for field, index in mapping.items()
        }
        for required in ("engineering_no", "section_no", "pallet_no", "pallet_name"):
            if not fields.get(required):
                diagnostics.append(
                    {
                        "code": "MISSING_REQUIRED_FIELD",
                        "excel_row": excel_row,
                        "field": required,
                        "message": f"{labels.get(required, required)}为空",
                    }
                )
        if any(
            not fields.get(field)
            for field in ("engineering_no", "section_no", "pallet_no", "pallet_name")
        ):
            continue
        key = (
            fields["engineering_no"].casefold(),
            fields["section_no"].casefold(),
            fields["pallet_no"].casefold(),
        )
        if key in seen_pallets:
            diagnostics.append(
                {
                    "code": "DUPLICATE_DEMAND_PALLET",
                    "excel_row": excel_row,
                    "engineering_no": fields["engineering_no"],
                    "section_no": fields["section_no"],
                    "pallet_no": fields["pallet_no"],
                    "message": "同一工作簿内工程号、分段号和托盘号重复",
                }
            )
            continue
        seen_pallets.add(key)

        demand_start_at, demand_end_at = _parse_window(fields.get("demand_window_raw"))
        if fields.get("demand_window_raw") and not (demand_start_at and demand_end_at):
            diagnostics.append(
                {
                    "code": "INVALID_DEMAND_WINDOW",
                    "excel_row": excel_row,
                    "value": fields.get("demand_window_raw"),
                    "message": "需求时间无法解析为起止时间",
                }
            )
            continue
        if demand_start_at and demand_end_at and demand_end_at < demand_start_at:
            diagnostics.append(
                {
                    "code": "INVALID_DEMAND_WINDOW",
                    "excel_row": excel_row,
                    "value": fields.get("demand_window_raw"),
                    "message": "需求结束时间早于开始时间",
                }
            )
            continue

        delivery_status = (
            _to_text(values[delivery_status_index])
            if delivery_status_index is not None and delivery_status_index < len(values)
            else None
        )
        raw_data = {}
        for column_index, value in enumerate(values):
            header = _to_text(rows[header_offset][1][column_index])
            key_name = header or f"未命名列{get_column_letter(column_index + 1)}"
            raw_data[key_name] = _json_value(value)
        raw_data["excel_row"] = excel_row
        raw_data["sheet_name"] = sheet_name
        raw_data["form_template"] = "shipyard_distribution_demand"

        line_payload = {
            **fields,
            "delivery_status": delivery_status,
            "engineering_no": fields["engineering_no"],
            "section_no": fields["section_no"],
            "pallet_no": fields["pallet_no"],
            "source_row_no": str(excel_row),
            "raw_data": raw_data,
        }
        lines.append(
            ParsedDistributionLine(
                source_row_no=str(excel_row),
                source_row_hash=_row_hash(line_payload),
                engineering_no=fields["engineering_no"],
                section_no=fields["section_no"],
                pallet_no=fields["pallet_no"],
                pallet_name=fields["pallet_name"],
                work_order_status=fields.get("work_order_status"),
                delivery_status=delivery_status,
                applicant_raw=fields.get("applicant_raw"),
                department=fields.get("department"),
                delivery_address=fields.get("delivery_address"),
                contact_name=fields.get("contact_name"),
                contact_phone=fields.get("contact_phone"),
                demand_window_raw=fields.get("demand_window_raw"),
                demand_start_at=demand_start_at,
                demand_end_at=demand_end_at,
                resource_1=fields.get("resource_1"),
                driver_1=fields.get("driver_1"),
                resource_2=fields.get("resource_2"),
                driver_2=fields.get("driver_2"),
                resource_3=fields.get("resource_3"),
                driver_3=fields.get("driver_3"),
                confirmer=fields.get("confirmer"),
                completed_at_raw=fields.get("completed_at_raw"),
                completed_at=_parse_datetime(fields.get("completed_at_raw")),
                confirmation_status=fields.get("confirmation_status"),
                raw_data=raw_data,
            )
        )
    return lines, diagnostics


def parse_distribution_workbook(
    content: bytes,
    *,
    source_file_name: str | None = None,
) -> ParsedDistributionWorkbook:
    if not content:
        raise DistributionParseError("上传文件为空")
    try:
        validate_xlsx_archive(content)
    except XlsxSecurityError as exc:
        raise DistributionParseError(
            str(exc),
            [{"code": exc.code}],
        ) from exc
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise DistributionParseError(
            "无法读取 Excel 文件，请确认文件未损坏且为 .xlsx 格式"
        ) from exc

    best = None
    best_score = -1
    sheet_rows = {}
    try:
        for worksheet in workbook.worksheets:
            rows = _read_sheet_rows(worksheet)
            sheet_rows[worksheet.title] = rows
            candidate = _find_header(rows)
            if candidate is None:
                continue
            score = len(candidate[2])
            if score > best_score:
                best_score = score
                best = (worksheet.title, rows, candidate)
    finally:
        workbook.close()

    if best is None:
        raise DistributionParseError(
            "未找到船厂托盘配送需求表头",
            [
                {
                    "code": "HEADER_NOT_FOUND",
                    "message": "至少需要工程号、分段号、托盘号和托盘名称",
                }
            ],
        )

    sheet_name, rows, (header_offset, header_row, mapping, labels) = best
    delivery_status_index, warnings = _delivery_status_column(
        rows,
        header_offset=header_offset,
        mapping=mapping,
    )
    lines, diagnostics = _business_rows(
        rows=rows,
        header_offset=header_offset,
        mapping=mapping,
        labels=labels,
        delivery_status_index=delivery_status_index,
        sheet_name=sheet_name,
    )
    if diagnostics:
        raise DistributionParseError("配送需求表存在无法导入的行", diagnostics)
    if not lines:
        raise DistributionParseError("配送需求表没有业务行")

    ignored_sheets = [
        name for name, sheet_rows in sheet_rows.items() if name != sheet_name
    ]
    metadata = {
        "parser_version": PARSER_VERSION,
        "source_file_name": source_file_name,
        "sheet_name": sheet_name,
        "header_row": header_row,
        "line_count": len(lines),
        "engineering_count": len({line.engineering_no for line in lines}),
        "pallet_count": len({line.pallet_no for line in lines}),
        "ignored_sheets": ignored_sheets,
        "field_mapping": {
            **labels,
            **(
                {
                    "delivery_status": (
                        f"未命名列{get_column_letter(delivery_status_index + 1)}"
                    )
                }
                if delivery_status_index is not None
                else {}
            ),
        },
        "source_timezone": "Asia/Shanghai",
    }
    return ParsedDistributionWorkbook(
        sheet_name=sheet_name,
        header_row=header_row,
        metadata=metadata,
        warnings=tuple(warnings),
        lines=tuple(lines),
    )
