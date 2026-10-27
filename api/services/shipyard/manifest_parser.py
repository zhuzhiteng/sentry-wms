"""Parser for the shipyard factory inbound manifest template.

The template is supplied by the factory. Its title, header row and footer
position are not treated as fixed cell coordinates: the parser finds the
header by normalized labels, then reads the business rows beneath it.

One source row is one cargo ticket / pallet group. The quantity column is the
quantity of the whole row, not a per-pallet quantity.
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, Optional
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


MAX_SCAN_ROWS = 5000
MAX_TEMPLATE_COLUMNS = 64
HEADER_SEARCH_ROWS = 30
MAX_ARCHIVE_ENTRIES = 2000
MAX_UNCOMPRESSED_ARCHIVE_BYTES = 50 * 1024 * 1024

HEADER_ALIASES = {
    "source_row_no": ("序号", "清单序号", "行号"),
    "ship_no": ("船号",),
    "engineering_no": ("工程号", "项目号"),
    "item_name": ("品名", "物料名称", "物资名称"),
    "production_drawing_no": ("图号", "制作图编号", "制作图号"),
    "drawing_name": ("图名", "制作图名称", "安装托盘名称"),
    "declared_quantity": ("件数", "数量", "整票件数", "物资数量"),
    "delivery_note_no": ("送货单号", "送货单编号"),
    "po_no": ("PO号", "采购订单号"),
    "po_reference": ("PO NO", "PO参考号"),
    "destination": ("送货地点", "目的地"),
    "buyer": ("采购员",),
    "weight": ("重量", "总重"),
    "cargo_category": ("螺丝", "货物类别", "备注"),
    "unit": ("单位",),
    "item_spec": ("型号规格", "规格型号", "型号"),
}


@dataclass(frozen=True)
class ParsedInboundLine:
    source_row_no: str
    source_row_hash: str
    ship_no: str
    engineering_no: str
    production_drawing_no: str
    drawing_name: Optional[str]
    manufacturer_name: str
    delivery_note_no: Optional[str]
    po_no: Optional[str]
    item_name: str
    item_spec: Optional[str]
    unit: str
    declared_quantity: Decimal
    raw_data: dict[str, Any]


@dataclass(frozen=True)
class ParsedInboundWorkbook:
    sheet_name: str
    header_row: int
    metadata: dict[str, Any]
    lines: tuple[ParsedInboundLine, ...]


class ManifestParseError(ValueError):
    """Raised when an uploaded workbook cannot be safely converted to lines."""

    def __init__(self, message: str, diagnostics: Optional[list[dict[str, Any]]] = None):
        super().__init__(message)
        self.diagnostics = diagnostics or []


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    return re.sub(r"[\s\u3000:：()（）]+", "", text)


_NORMALIZED_ALIASES = {
    canonical: frozenset(_normalize_header(alias) for alias in aliases)
    for canonical, aliases in HEADER_ALIASES.items()
}


def _to_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, Decimal):
        text = format(value, "f")
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = text.strip()
    return text or None


def _to_decimal(value: Any) -> Optional[Decimal]:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", "").replace("，", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _cell(values: tuple[Any, ...], index: Optional[int]) -> Any:
    if index is None or index >= len(values):
        return None
    return values[index]


def _find_header(
    rows: list[tuple[int, tuple[Any, ...]]],
) -> Optional[tuple[int, int, dict[str, int], dict[str, str]]]:
    best = None
    best_score = -1
    for row_offset, (excel_row, values) in enumerate(rows[:HEADER_SEARCH_ROWS]):
        mapping: dict[str, int] = {}
        labels: dict[str, str] = {}
        for column_index, value in enumerate(values[:MAX_TEMPLATE_COLUMNS]):
            normalized = _normalize_header(value)
            if not normalized:
                continue
            for canonical, aliases in _NORMALIZED_ALIASES.items():
                if canonical not in mapping and normalized in aliases:
                    mapping[canonical] = column_index
                    labels[canonical] = str(value).strip()
                    break

        has_core_fields = (
            "ship_no" in mapping
            and "declared_quantity" in mapping
            and ({"item_name", "drawing_name", "production_drawing_no"} & mapping.keys())
        )
        if not has_core_fields:
            continue
        score = len(mapping)
        if score > best_score:
            best_score = score
            best = (row_offset, excel_row, mapping, labels)
    return best


def _read_sheet_rows(worksheet) -> list[tuple[int, tuple[Any, ...]]]:
    max_row = min(worksheet.max_row or 0, MAX_SCAN_ROWS)
    if max_row <= 0:
        return []
    max_column = min(worksheet.max_column or 1, MAX_TEMPLATE_COLUMNS)
    return [
        (excel_row, tuple(values))
        for excel_row, values in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=max_row,
                max_col=max_column,
                values_only=True,
            ),
            start=1,
        )
    ]


def _extract_title_rows(
    rows: list[tuple[int, tuple[Any, ...]]], header_offset: int
) -> list[str]:
    titles: list[str] = []
    for _, values in rows[:header_offset]:
        for value in values:
            text = _to_text(value)
            if text and text not in titles:
                titles.append(text)
    return titles


def _extract_manufacturer(title_rows: list[str]) -> str:
    for text in title_rows:
        if any(marker in text for marker in ("公司", "船厂", "工厂", "厂")):
            return text[:200]
    return ""


def _extract_document_date(rows: list[tuple[int, tuple[Any, ...]]]) -> Optional[str]:
    for _, values in rows:
        for value in values:
            if isinstance(value, datetime):
                return value.date().isoformat()
            if isinstance(value, date):
                return value.isoformat()
            text = _to_text(value)
            if text and re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
                return text
    return None


def _row_hash(payload: dict[str, Any]) -> str:
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise ManifestParseError(
                    "Excel 文件包含过多内部文件，拒绝解析",
                    [{"code": "XLSX_TOO_MANY_ENTRIES"}],
                )
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_ARCHIVE_BYTES:
                raise ManifestParseError(
                    "Excel 解压后体积过大，拒绝解析",
                    [{"code": "XLSX_UNCOMPRESSED_TOO_LARGE"}],
                )
    except BadZipFile as exc:
        raise ManifestParseError(
            "无法读取 Excel 文件，请确认文件为有效的 .xlsx 文件"
        ) from exc


def parse_inbound_workbook(
    content: bytes,
    *,
    manufacturer_name: Optional[str] = None,
    source_file_name: Optional[str] = None,
) -> ParsedInboundWorkbook:
    if not content:
        raise ManifestParseError("上传文件为空")
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise ManifestParseError("无法读取 Excel 文件，请确认文件未损坏且为 .xlsx 格式") from exc

    best = None
    best_score = -1
    sheet_rows_by_name: dict[str, list[tuple[int, tuple[Any, ...]]]] = {}
    try:
        for worksheet in workbook.worksheets:
            rows = _read_sheet_rows(worksheet)
            sheet_rows_by_name[worksheet.title] = rows
            candidate = _find_header(rows)
            if candidate is None:
                continue
            score = len(candidate[2])
            if score > best_score:
                best_score = score
                best = (worksheet.title, rows, candidate)
    finally:
        workbook.close()

    if best is None:
        raise ManifestParseError(
            "未找到发货清单表头",
            [
                {
                    "code": "HEADER_NOT_FOUND",
                    "message": "至少需要包含船号、件数，以及品名/图号/图名之一",
                }
            ],
        )

    sheet_name, rows, (header_offset, header_excel_row, mapping, labels) = best
    title_rows = _extract_title_rows(rows, header_offset)
    resolved_manufacturer = (
        (manufacturer_name or "").strip()[:200]
        or _extract_manufacturer(title_rows)
    )

    lines: list[ParsedInboundLine] = []
    diagnostics: list[dict[str, Any]] = []
    seen_source_rows: set[str] = set()

    for excel_row, values in rows[header_offset + 1 :]:
        ship_no = _to_text(_cell(values, mapping.get("ship_no")))
        quantity_raw = _cell(values, mapping.get("declared_quantity"))
        item_name = _to_text(_cell(values, mapping.get("item_name")))
        drawing_name = _to_text(_cell(values, mapping.get("drawing_name")))
        drawing_no = _to_text(_cell(values, mapping.get("production_drawing_no")))
        serial = _to_text(_cell(values, mapping.get("source_row_no")))

        non_empty_values = [value for value in values if _to_text(value) is not None]
        if not non_empty_values:
            continue
        if not ship_no:
            # Footer, notes and unrelated auxiliary blocks are not business rows.
            if item_name or drawing_name or drawing_no or quantity_raw is not None:
                diagnostics.append(
                    {
                        "code": "MISSING_SHIP_NO",
                        "excel_row": excel_row,
                        "message": "存在货物数据但船号为空",
                    }
                )
            continue

        if quantity_raw is None or _to_text(quantity_raw) is None:
            diagnostics.append(
                {
                    "code": "MISSING_QUANTITY",
                    "excel_row": excel_row,
                    "ship_no": ship_no,
                    "message": "件数为空",
                }
            )
            continue

        declared_quantity = _to_decimal(quantity_raw)
        if declared_quantity is None or declared_quantity <= 0:
            diagnostics.append(
                {
                    "code": "INVALID_QUANTITY",
                    "excel_row": excel_row,
                    "ship_no": ship_no,
                    "value": _json_value(quantity_raw),
                    "message": "件数必须为正数",
                }
            )
            continue

        resolved_item_name = item_name or drawing_name or drawing_no
        if not resolved_item_name:
            diagnostics.append(
                {
                    "code": "MISSING_ITEM_NAME",
                    "excel_row": excel_row,
                    "ship_no": ship_no,
                    "message": "品名、图名和图号均为空",
                }
            )
            continue

        source_row_no = serial or str(excel_row)
        if source_row_no in seen_source_rows:
            diagnostics.append(
                {
                    "code": "DUPLICATE_SOURCE_ROW",
                    "excel_row": excel_row,
                    "source_row_no": source_row_no,
                    "message": "清单行号重复，无法建立一一对应关系",
                }
            )
            continue
        seen_source_rows.add(source_row_no)

        raw_data = {
            labels[canonical]: _json_value(_cell(values, column_index))
            for canonical, column_index in mapping.items()
        }
        raw_data.update(
            {
                "excel_row": excel_row,
                "sheet_name": sheet_name,
                "form_template": "factory_inbound_manifest",
            }
        )
        line_payload = {
            "source_row_no": source_row_no,
            "ship_no": ship_no,
            "production_drawing_no": drawing_no or "",
            "drawing_name": drawing_name or "",
            "declared_quantity": format(declared_quantity, "f"),
        }
        lines.append(
            ParsedInboundLine(
                source_row_no=source_row_no,
                source_row_hash=_row_hash(line_payload),
                ship_no=ship_no,
                engineering_no=(
                    _to_text(_cell(values, mapping.get("engineering_no"))) or ""
                ),
                production_drawing_no=drawing_no or "",
                drawing_name=drawing_name,
                manufacturer_name=resolved_manufacturer,
                delivery_note_no=_to_text(
                    _cell(values, mapping.get("delivery_note_no"))
                ),
                po_no=(
                    _to_text(_cell(values, mapping.get("po_no")))
                    or _to_text(_cell(values, mapping.get("po_reference")))
                ),
                item_name=resolved_item_name[:255],
                item_spec=_to_text(_cell(values, mapping.get("item_spec"))),
                unit=_to_text(_cell(values, mapping.get("unit"))) or "件",
                declared_quantity=declared_quantity,
                raw_data=raw_data,
            )
        )

    if diagnostics:
        raise ManifestParseError("清单存在无法导入的行", diagnostics)
    if not lines:
        raise ManifestParseError(
            "清单中没有可导入的业务行",
            [{"code": "NO_BUSINESS_ROWS", "message": "表头下未找到有效货物行"}],
        )

    ignored_sheets = [
        name for name, sheet_rows in sheet_rows_by_name.items() if name != sheet_name
    ]
    total_quantity = sum((line.declared_quantity for line in lines), Decimal("0"))
    metadata = {
        "parser_version": "shipyard-excel-v1",
        "source_file_name": source_file_name,
        "sheet_name": sheet_name,
        "header_row": header_excel_row,
        "header_labels": labels,
        "title_rows": title_rows,
        "manufacturer_name": resolved_manufacturer,
        "document_date": _extract_document_date(rows),
        "ignored_sheets": ignored_sheets,
        "line_count": len(lines),
        "ship_count": len({line.ship_no for line in lines}),
        "total_quantity": format(total_quantity, "f"),
        "row_contract": "one_source_row_equals_one_ticket_or_ticket_group",
        "quantity_contract": "row_quantity_is_whole_ticket_group_total",
    }
    return ParsedInboundWorkbook(
        sheet_name=sheet_name,
        header_row=header_excel_row,
        metadata=metadata,
        lines=tuple(lines),
    )
