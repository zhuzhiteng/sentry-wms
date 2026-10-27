"""Parser for the shipyard formal outbound manifest.

One source row is one external demand for a complete cargo ticket or ticket
group. Material code and item name are required because stock matching must
never fall back to a name-only heuristic.
"""

from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from services.shipyard.xlsx_security import XlsxSecurityError, validate_xlsx_archive

MAX_SCAN_ROWS = 5000
MAX_TEMPLATE_COLUMNS = 80
HEADER_SEARCH_ROWS = 40

HEADER_ALIASES = {
    "source_row_no": ("清单行号", "行号", "序号", "票号", "需求行号"),
    "factory_shipment_no": ("船厂出货单号", "出货单号", "发货单号"),
    "ticket_no": ("票货号", "托组号", "票货/托组", "清单票号"),
    "ship_no": ("船号",),
    "engineering_no": ("工程号", "项目号"),
    "production_drawing_no": ("制作图编号", "制作图号", "图号"),
    "manufacturer_name": ("厂家", "制造厂家", "供应商"),
    "material_code": ("物料编码", "物资编码", "物料代码", "物资代码"),
    "item_name": ("物料名称", "品名", "物资名称"),
    "item_spec": ("规格型号", "型号规格", "物资规格", "规格"),
    "demand_quantity": ("需求数量", "出库数量", "数量", "件数"),
    "unit": ("单位",),
    "demand_date": ("需求日期", "出库日期", "需求时间"),
    "expected_pallet_count": ("预计托盘数", "预计托数", "托数"),
    "expected_tag_suffixes": ("托盘后缀", "吊牌后缀", "托盘清单", "托组后缀"),
}

REQUIRED_COLUMNS = {
    "source_row_no",
    "ship_no",
    "engineering_no",
    "production_drawing_no",
    "material_code",
    "item_name",
    "demand_quantity",
}


@dataclass(frozen=True)
class ParsedOutboundLine:
    source_row_no: str
    ticket_no: str | None
    ship_no: str
    engineering_no: str
    production_drawing_no: str
    manufacturer_name: str
    material_code: str
    item_name: str
    item_spec: str | None
    demand_quantity: Decimal
    unit: str
    demand_date: date | None
    expected_pallet_count: int | None
    expected_tag_suffixes: tuple[str, ...]
    raw_data: dict[str, Any]


@dataclass(frozen=True)
class ParsedOutboundWorkbook:
    sheet_name: str
    header_row: int
    factory_shipment_no: str | None
    metadata: dict[str, Any]
    lines: tuple[ParsedOutboundLine, ...]


class OutboundManifestParseError(ValueError):
    def __init__(self, message: str, diagnostics: list[dict[str, Any]] | None = None):
        super().__init__(message)
        self.diagnostics = diagnostics or []


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    return re.sub(r"[\s\u3000:：()（）/]+", "", text)


_NORMALIZED_ALIASES = {
    canonical: frozenset(_normalize_header(alias) for alias in aliases)
    for canonical, aliases in HEADER_ALIASES.items()
}


def _to_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, Decimal):
        text = format(value, "f")
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = text.strip()
    return text or None


def _to_decimal(value: Any) -> Decimal | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, (Decimal, int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", "").replace("，", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _to_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _to_text(value)
    if not text:
        return None
    match = re.match(r"^(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})", text)
    if match:
        try:
            return date(
                int(match.group(1)),
                int(match.group(2)),
                int(match.group(3)),
            )
        except ValueError:
            return None
    return None


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _row_hash(payload: dict[str, Any]) -> str:
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _cell(values: tuple[Any, ...], index: int | None) -> Any:
    if index is None or index >= len(values):
        return None
    return values[index]


def _read_sheet_rows(worksheet) -> list[tuple[int, tuple[Any, ...]]]:
    max_row = min(worksheet.max_row or 0, MAX_SCAN_ROWS)
    if max_row <= 0:
        return []
    max_column = min(worksheet.max_column or 1, MAX_TEMPLATE_COLUMNS)
    return [
        (excel_row, tuple(values))
        for excel_row, values in enumerate(
            worksheet.iter_rows(
                min_row=1,
                max_row=max_row,
                max_col=max_column,
                values_only=True,
            ),
            start=1,
        )
    ]


def _find_header(
    rows: list[tuple[int, tuple[Any, ...]]],
) -> tuple[int, int, dict[str, int], dict[str, str]] | None:
    best = None
    best_score = -1
    for row_offset, (excel_row, values) in enumerate(rows[:HEADER_SEARCH_ROWS]):
        mapping: dict[str, int] = {}
        labels: dict[str, str] = {}
        for column_index, value in enumerate(values[:MAX_TEMPLATE_COLUMNS]):
            normalized = _normalize_header(value)
            if not normalized:
                continue
            for canonical, aliases in _NORMALIZED_ALIASES.items():
                if canonical not in mapping and normalized in aliases:
                    mapping[canonical] = column_index
                    labels[canonical] = str(value).strip()
                    break
        if not REQUIRED_COLUMNS.issubset(mapping):
            continue
        score = len(mapping)
        if score > best_score:
            best_score = score
            best = (row_offset, excel_row, mapping, labels)
    return best


def _extract_title_rows(
    rows: list[tuple[int, tuple[Any, ...]]], header_offset: int
) -> list[str]:
    titles: list[str] = []
    for _, values in rows[:header_offset]:
        for value in values:
            text = _to_text(value)
            if text and text not in titles:
                titles.append(text)
    return titles


def _extract_document_date(rows: list[tuple[int, tuple[Any, ...]]]) -> date | None:
    for _, values in rows:
        for value in values:
            parsed = _to_date(value)
            if parsed is not None:
                return parsed
    return None


def _parse_suffixes(value: Any) -> tuple[list[str], str | None]:
    text = _to_text(value)
    if not text:
        return [], None
    parts = [item.strip() for item in re.split(r"[,，、;；\s]+", text) if item.strip()]
    invalid = [item for item in parts if not re.fullmatch(r"^[1-9]\d*-[1-9]\d*$", item)]
    if invalid:
        return [], f"托盘后缀格式无效：{', '.join(invalid)}"
    if len(set(parts)) != len(parts):
        return [], "托盘后缀存在重复"
    return parts, None


def _parse_business_row(
    *,
    excel_row: int,
    values: tuple[Any, ...],
    mapping: dict[str, int],
    labels: dict[str, str],
    sheet_name: str,
    seen_source_rows: set[str],
) -> tuple[ParsedOutboundLine | None, list[dict[str, Any]]]:
    if not any(_to_text(value) is not None for value in values):
        return None, []

    source_row_no = _to_text(_cell(values, mapping["source_row_no"]))
    if not source_row_no:
        return None, [
            {
                "code": "MISSING_SOURCE_ROW",
                "excel_row": excel_row,
                "message": "清单行号为空",
            }
        ]
    if source_row_no in seen_source_rows:
        return None, [
            {
                "code": "DUPLICATE_SOURCE_ROW",
                "excel_row": excel_row,
                "source_row_no": source_row_no,
                "message": "清单行号重复",
            }
        ]

    required_text = {
        "ship_no": _to_text(_cell(values, mapping["ship_no"])),
        "engineering_no": _to_text(_cell(values, mapping["engineering_no"])),
        "production_drawing_no": _to_text(
            _cell(values, mapping["production_drawing_no"])
        ),
        "material_code": _to_text(_cell(values, mapping["material_code"])),
        "item_name": _to_text(_cell(values, mapping["item_name"])),
    }
    missing = [field for field, text in required_text.items() if not text]
    if missing:
        return None, [
            {
                "code": "MISSING_REQUIRED_FIELDS",
                "excel_row": excel_row,
                "source_row_no": source_row_no,
                "fields": missing,
                "message": "船号、工程号、制作图、物料编码和物料名称均为必填",
            }
        ]

    demand_quantity = _to_decimal(_cell(values, mapping["demand_quantity"]))
    if demand_quantity is None or demand_quantity <= 0:
        return None, [
            {
                "code": "INVALID_DEMAND_QUANTITY",
                "excel_row": excel_row,
                "source_row_no": source_row_no,
                "message": "需求数量必须为正数",
            }
        ]

    suffix_text = _cell(values, mapping.get("expected_tag_suffixes"))
    suffixes, suffix_error = _parse_suffixes(suffix_text)
    if suffix_error:
        return None, [
            {
                "code": "INVALID_TAG_SUFFIXES",
                "excel_row": excel_row,
                "source_row_no": source_row_no,
                "message": suffix_error,
            }
        ]

    count_raw = _cell(values, mapping.get("expected_pallet_count"))
    expected_pallet_count = None
    if _to_text(count_raw) is not None:
        count_decimal = _to_decimal(count_raw)
        if (
            count_decimal is None
            or count_decimal <= 0
            or count_decimal != count_decimal.to_integral_value()
        ):
            return None, [
                {
                    "code": "INVALID_EXPECTED_PALLET_COUNT",
                    "excel_row": excel_row,
                    "source_row_no": source_row_no,
                    "message": "预计托盘数必须为正整数",
                }
            ]
        expected_pallet_count = int(count_decimal)
    if suffixes and expected_pallet_count is None:
        expected_pallet_count = len(suffixes)
    if (
        suffixes
        and expected_pallet_count is not None
        and len(suffixes) != expected_pallet_count
    ):
        return None, [
            {
                "code": "PALLET_SUFFIX_COUNT_MISMATCH",
                "excel_row": excel_row,
                "source_row_no": source_row_no,
                "expected_pallet_count": expected_pallet_count,
                "suffix_count": len(suffixes),
                "message": "托盘后缀数量与预计托盘数不一致",
            }
        ]

    date_raw = _cell(values, mapping.get("demand_date"))
    demand_date = _to_date(date_raw) if _to_text(date_raw) else None
    if _to_text(date_raw) and demand_date is None:
        return None, [
            {
                "code": "INVALID_DEMAND_DATE",
                "excel_row": excel_row,
                "source_row_no": source_row_no,
                "message": "需求日期无法识别",
            }
        ]

    seen_source_rows.add(source_row_no)
    material_code = required_text["material_code"][:128]
    item_name = required_text["item_name"][:255]
    line_payload = {
        "source_row_no": source_row_no,
        "ship_no": required_text["ship_no"],
        "engineering_no": required_text["engineering_no"],
        "production_drawing_no": required_text["production_drawing_no"],
        "material_code": material_code,
        "item_name": item_name,
        "item_spec": _to_text(_cell(values, mapping.get("item_spec"))),
        "demand_quantity": format(demand_quantity, "f"),
        "expected_tag_suffixes": suffixes,
    }
    raw_data = {
        labels[canonical]: _json_value(_cell(values, column_index))
        for canonical, column_index in mapping.items()
    }
    raw_data.update(
        {
            "excel_row": excel_row,
            "sheet_name": sheet_name,
            "form_template": "shipyard_formal_outbound",
            "source_row_hash": _row_hash(line_payload),
        }
    )
    return (
        ParsedOutboundLine(
            source_row_no=source_row_no,
            ticket_no=_to_text(_cell(values, mapping.get("ticket_no"))),
            ship_no=required_text["ship_no"],
            engineering_no=required_text["engineering_no"],
            production_drawing_no=required_text["production_drawing_no"],
            manufacturer_name=(
                _to_text(_cell(values, mapping.get("manufacturer_name"))) or ""
            )[:200],
            material_code=material_code,
            item_name=item_name,
            item_spec=_to_text(_cell(values, mapping.get("item_spec"))),
            demand_quantity=demand_quantity,
            unit=_to_text(_cell(values, mapping.get("unit"))) or "件",
            demand_date=demand_date,
            expected_pallet_count=expected_pallet_count,
            expected_tag_suffixes=tuple(suffixes),
            raw_data=raw_data,
        ),
        [],
    )


def parse_outbound_workbook(
    content: bytes,
    *,
    source_file_name: str | None = None,
) -> ParsedOutboundWorkbook:
    if not content:
        raise OutboundManifestParseError("上传文件为空")
    try:
        validate_xlsx_archive(content)
    except XlsxSecurityError as exc:
        raise OutboundManifestParseError(
            str(exc),
            [{"code": exc.code}],
        ) from exc
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise OutboundManifestParseError(
            "无法读取 Excel 文件，请确认文件未损坏且为 .xlsx 格式"
        ) from exc

    best = None
    best_score = -1
    sheet_rows_by_name: dict[str, list[tuple[int, tuple[Any, ...]]]] = {}
    try:
        for worksheet in workbook.worksheets:
            rows = _read_sheet_rows(worksheet)
            sheet_rows_by_name[worksheet.title] = rows
            candidate = _find_header(rows)
            if candidate is None:
                continue
            score = len(candidate[2])
            if score > best_score:
                best_score = score
                best = (worksheet.title, rows, candidate)
    finally:
        workbook.close()

    if best is None:
        missing_titles = "、".join(
            alias[0] for alias in (HEADER_ALIASES[key] for key in REQUIRED_COLUMNS)
        )
        raise OutboundManifestParseError(
            "未找到正式出库清单表头",
            [
                {
                    "code": "HEADER_NOT_FOUND",
                    "message": f"固定模板至少需要：{missing_titles}",
                }
            ],
        )

    sheet_name, rows, (header_offset, header_excel_row, mapping, labels) = best
    lines: list[ParsedOutboundLine] = []
    diagnostics: list[dict[str, Any]] = []
    seen_source_rows: set[str] = set()
    for excel_row, values in rows[header_offset + 1 :]:
        line, row_diagnostics = _parse_business_row(
            excel_row=excel_row,
            values=values,
            mapping=mapping,
            labels=labels,
            sheet_name=sheet_name,
            seen_source_rows=seen_source_rows,
        )
        diagnostics.extend(row_diagnostics)
        if line is not None:
            lines.append(line)

    if diagnostics:
        raise OutboundManifestParseError("出库清单存在无法导入的行", diagnostics)
    if not lines:
        raise OutboundManifestParseError(
            "出库清单中没有有效业务行",
            [{"code": "NO_BUSINESS_ROWS", "message": "表头下未找到有效出库行"}],
        )

    factory_shipment_nos = {
        _to_text(_cell(values, mapping.get("factory_shipment_no")))
        for _, values in rows[header_offset + 1 :]
    }
    factory_shipment_nos.discard(None)
    if len(factory_shipment_nos) > 1:
        raise OutboundManifestParseError(
            "出库清单存在多个船厂出货单号",
            [
                {
                    "code": "MULTIPLE_FACTORY_SHIPMENT_NOS",
                    "values": sorted(factory_shipment_nos),
                    "message": "一个导入文件只能包含一个船厂出货单号",
                }
            ],
        )

    title_rows = _extract_title_rows(rows, header_offset)
    ignored_sheets = [
        name for name, sheet_rows in sheet_rows_by_name.items() if name != sheet_name
    ]
    total_quantity = sum((line.demand_quantity for line in lines), Decimal(0))
    document_date = _extract_document_date(rows)
    metadata = {
        "parser_version": "shipyard-outbound-v1",
        "source_file_name": source_file_name,
        "sheet_name": sheet_name,
        "header_row": header_excel_row,
        "header_labels": labels,
        "title_rows": title_rows,
        "document_date": document_date.isoformat() if document_date else None,
        "ignored_sheets": ignored_sheets,
        "line_count": len(lines),
        "ship_count": len({line.ship_no for line in lines}),
        "total_quantity": format(total_quantity, "f"),
        "row_contract": "one_source_row_is_one_external_outbound_demand",
        "quantity_contract": "quantity_is_whole_ticket_or_ticket_group_quantity",
        "material_contract": "material_code_and_item_name_are_required",
    }
    return ParsedOutboundWorkbook(
        sheet_name=sheet_name,
        header_row=header_excel_row,
        factory_shipment_no=(
            next(iter(factory_shipment_nos)) if factory_shipment_nos else None
        ),
        metadata=metadata,
        lines=tuple(lines),
    )
